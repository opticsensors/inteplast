"""Lectores por formato. Cada uno recibe la ruta larga del original y devuelve un Resultado.

Reglas comunes:
- Solo lectura. Lo que necesita una aplicacion (Word/Excel por COM) se abre sobre una copia
  temporal, para que nunca aparezcan ficheros de bloqueo (~$) dentro de Exemples.
- Un valor ausente no es cero: las celdas vacias no se guardan.
- Si un fichero es demasiado grande se recorta y se dice en `mensaje`/`estructura`.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

MAX_CELDAS = 400_000        # por libro
MAX_LINEAS_TEXTO = 20_000   # por fichero de texto
SIETE_ZIP = r"C:\Program Files\7-Zip\7z.exe"


@dataclass
class Resultado:
    estado: str = "extraido"   # extraido | solo_metadatos | sin_texto | error
    elementos: list[tuple[str, str, str, dict | None]] = field(default_factory=list)
    celdas: list[tuple[str, int, int, str, float | None, str | None]] = field(default_factory=list)
    estructura: dict[str, Any] = field(default_factory=dict)
    mensaje: str | None = None

    def elemento(self, tipo: str, localizador: str, texto: str, datos: dict | None = None) -> None:
        self.elementos.append((tipo, localizador, texto, datos))


# ---------------------------------------------------------------------------------------
# utilidades
# ---------------------------------------------------------------------------------------


def numero(valor: Any) -> float | None:
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor) if math.isfinite(valor) else None
    try:
        x = float(str(valor).strip().replace(",", "."))
        return x if math.isfinite(x) else None
    except (TypeError, ValueError):
        return None


def texto_valor(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float):
        return f"{valor:.15g}"
    if isinstance(valor, (dt.datetime, dt.date, dt.time)):
        return valor.isoformat()
    return str(valor)


def decodificar(data: bytes) -> tuple[str, str]:
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(codificacion), codificacion
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "latin-1"


def columna(indice: int) -> str:
    """1 -> A, 27 -> AA."""
    letras = ""
    while indice:
        indice, resto = divmod(indice - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


def cadenas_imprimibles(data: bytes, minimo: int = 6, limite: int = 400) -> list[str]:
    """Cadenas ASCII y UTF-16LE legibles dentro de un binario (CATIA, Moldflow, desconocidos)."""
    vistas: dict[str, None] = {}
    for patron in (rb"[\x20-\x7e]{%d,}" % minimo, rb"(?:[\x20-\x7e]\x00){%d,}" % minimo):
        for m in re.finditer(patron, data):
            s = m.group().decode("utf-16-le" if b"\x00" in m.group() else "ascii", "ignore").strip()
            if len(s) >= minimo and not re.fullmatch(r"[\W_]+|(.)\1+", s):
                vistas.setdefault(s, None)
            if len(vistas) >= limite:
                return list(vistas)
    return list(vistas)


def copia_temporal(path: str) -> tuple[str, str]:
    carpeta = tempfile.mkdtemp(prefix="inteplast-lectura-")
    destino = os.path.join(carpeta, "copia" + Path(path).suffix)
    shutil.copyfile(path, destino)
    return carpeta, destino


def filas_a_elementos(r: Resultado, hoja: str, filas: dict[int, list[tuple[int, str]]]) -> None:
    for fila in sorted(filas):
        partes = [f"{columna(c)}{fila}={v}" for c, v in sorted(filas[fila]) if v.strip()]
        if partes:
            r.elemento("fila", f"{hoja}!{fila}", " | ".join(partes))


# ---------------------------------------------------------------------------------------
# tablas: CSV, XLS, XLSX/XLSM, ODS
# ---------------------------------------------------------------------------------------


def leer_csv(path: str) -> Resultado:
    r = Resultado()
    texto, codificacion = decodificar(Path(path).read_bytes())
    lineas = texto.splitlines()
    muestra = [ln for ln in lineas[:200] if ln.strip()]
    delimitador = max([";", ",", "\t"], key=lambda d: sum(ln.count(d) for ln in muestra))
    filas: dict[int, list[tuple[int, str]]] = {}
    numericas = 0
    for i, campos in enumerate(csv.reader(lineas, delimiter=delimitador), 1):
        valores = [(j, v) for j, v in enumerate(campos, 1) if v.strip()]
        if valores and all(numero(v) is not None for _, v in valores):
            numericas += 1
        if len(r.celdas) < MAX_CELDAS:
            for j, v in valores:
                r.celdas.append(("csv", i, j, v, numero(v), None))
        if valores and i <= MAX_LINEAS_TEXTO:
            filas[i] = valores
    filas_a_elementos(r, "csv", filas)
    bloques = sum(1 for ln in lineas if ln.startswith("****"))
    no_vacias = sum(1 for ln in lineas if ln.strip())
    r.estructura = {
        "codificacion": codificacion,
        "delimitador": delimitador,
        "lineas": len(lineas),
        "columnas_max": max((len(ln.split(delimitador)) for ln in muestra), default=0),
        "separadores_bloque": bloques,
        "formato": "cmm_bloques" if bloques else
                   "numerico" if no_vacias and numericas / no_vacias > 0.9 else "tabla",
        "primeras_lineas": muestra[:6],
    }
    if len(r.celdas) >= MAX_CELDAS:
        r.mensaje = f"Recortado a {MAX_CELDAS} celdas."
    return r


def propiedades_ole(path: str) -> dict[str, str]:
    import olefile

    try:
        with olefile.OleFileIO(path) as ole:
            meta = ole.get_metadata()
            campos = ["title", "subject", "author", "last_saved_by", "company",
                      "create_time", "last_saved_time", "creating_application"]
            return {c: texto_valor(getattr(meta, c, None).decode("cp1252", "ignore")
                                   if isinstance(getattr(meta, c, None), bytes)
                                   else getattr(meta, c, None))
                    for c in campos if getattr(meta, c, None)}
    except Exception:
        return {}


def valores_guardados(path: str) -> dict[str, list[list[Any]]]:
    """Valores tal como estan guardados en el fichero (calamine no recalcula). Vacio si no se puede."""
    try:
        from python_calamine import CalamineWorkbook

        libro = CalamineWorkbook.from_path(path)
        return {nombre: libro.get_sheet_by_name(nombre).to_python(skip_empty_area=False)
                for nombre in libro.sheet_names}
    except Exception:
        return {}


def xls_por_excel(path: str) -> Resultado:
    """XLS que xlrd no entiende: formulas con Excel por COM, sobre una copia; valores con calamine.

    Excel recalcula al abrir las funciones volatiles (RAND, RANDBETWEEN) y todo lo que depende de
    ellas: el valor de COM no es el guardado. Comprobado el 2026-10-21 en 13 libros (8.422 celdas
    distintas, todas RAND o dependientes). Si calamine no puede, se usa el valor de COM.
    """
    import pythoncom
    import win32com.client

    guardados = valores_guardados(path)
    r = Resultado(mensaje="Leido con Excel por COM (xlrd no admite este XLS)"
                  + ("; valores guardados leidos con calamine." if guardados else "; valores de COM (recalculados)."))
    carpeta, copia = copia_temporal(path)
    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    hojas = []
    try:
        libro = excel.Workbooks.Open(copia, 0, True)  # UpdateLinks=0, ReadOnly
        try:
            for hoja in libro.Worksheets:
                rango = hoja.UsedRange
                fila0, col0 = rango.Row, rango.Column
                valores, formulas = rango.Value, rango.Formula
                if not isinstance(valores, tuple):  # una sola celda
                    valores, formulas = ((valores,),), ((formulas,),)
                filas: dict[int, list[tuple[int, str]]] = {}
                n = 0
                rejilla = guardados.get(hoja.Name)
                for i, fila in enumerate(valores):
                    for j, valor in enumerate(fila):
                        if rejilla is not None:
                            f_g, c_g = fila0 + i - 1, col0 + j - 1
                            valor = rejilla[f_g][c_g] if f_g < len(rejilla) and c_g < len(rejilla[f_g]) else None
                            if isinstance(valor, float) and valor.is_integer() and isinstance(fila[j], int):
                                valor = int(valor)
                        texto = texto_valor(valor).strip()
                        if not texto:
                            continue
                        n += 1
                        if len(r.celdas) >= MAX_CELDAS:
                            continue
                        formula = formulas[i][j] if isinstance(formulas[i][j], str) and formulas[i][j].startswith("=") else None
                        r.celdas.append((hoja.Name, fila0 + i, col0 + j, texto,
                                         numero(valor) if not isinstance(valor, str) else None, formula))
                        filas.setdefault(fila0 + i, []).append((col0 + j, texto))
                filas_a_elementos(r, hoja.Name, filas)
                hojas.append({"nombre": hoja.Name, "filas": rango.Rows.Count, "columnas": rango.Columns.Count, "celdas": n})
        finally:
            libro.Close(False)
    finally:
        excel.Quit()
        pythoncom.CoUninitialize()
        shutil.rmtree(carpeta, ignore_errors=True)
    r.estructura = {"hojas": hojas, "propiedades": propiedades_ole(path), "formulas": "via Excel COM"}
    return r


def leer_xls(path: str) -> Resultado:
    import xlrd

    try:
        return _leer_xls(path)
    except xlrd.biffh.XLRDError as error:
        if "xlsx" in str(error).lower() or zipfile.is_zipfile(path):
            return leer_xlsx(path)
        return xls_por_excel(path)
    except AssertionError:
        return xls_por_excel(path)


def _leer_xls(path: str) -> Resultado:
    import xlrd

    r = Resultado()
    libro = xlrd.open_workbook(path, on_demand=True)
    hojas = []
    for hoja in libro.sheets():
        filas: dict[int, list[tuple[int, str]]] = {}
        n = 0
        for i in range(hoja.nrows):
            for j in range(hoja.ncols):
                celda = hoja.cell(i, j)
                if celda.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                if celda.ctype == xlrd.XL_CELL_DATE:
                    try:
                        valor = xlrd.xldate_as_datetime(celda.value, libro.datemode).isoformat()
                    except Exception:
                        valor = texto_valor(celda.value)
                elif celda.ctype == xlrd.XL_CELL_ERROR:
                    valor = xlrd.error_text_from_code.get(celda.value, "#ERROR")
                elif celda.ctype == xlrd.XL_CELL_BOOLEAN:
                    valor = "TRUE" if celda.value else "FALSE"
                else:
                    valor = texto_valor(celda.value).strip()
                if not valor:
                    continue
                n += 1
                if len(r.celdas) < MAX_CELDAS:
                    r.celdas.append((hoja.name, i + 1, j + 1, valor,
                                     numero(celda.value) if celda.ctype == xlrd.XL_CELL_NUMBER else None,
                                     None))
                    filas.setdefault(i + 1, []).append((j + 1, valor))
        filas_a_elementos(r, hoja.name, filas)
        hojas.append({"nombre": hoja.name, "filas": hoja.nrows, "columnas": hoja.ncols, "celdas": n})
        libro.unload_sheet(hoja.name)
    r.estructura = {"hojas": hojas, "propiedades": propiedades_ole(path),
                    "guardado_por": libro.user_name, "formulas": "no disponibles en XLS BIFF"}
    if len(r.celdas) >= MAX_CELDAS:
        r.mensaje = f"Recortado a {MAX_CELDAS} celdas."
    return r


def leer_xlsx(path: str) -> Resultado:
    import openpyxl

    r = Resultado()
    valores = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas: dict[tuple[str, int, int], str] = {}
    if os.path.getsize(path) < 25 * 1024 * 1024:
        crudo = openpyxl.load_workbook(path, read_only=True, data_only=False)
        for hoja in crudo.worksheets:
            for fila in hoja.iter_rows():
                for celda in fila:
                    if isinstance(celda.value, str) and celda.value.startswith("="):
                        formulas[(hoja.title, celda.row, celda.column)] = celda.value
        crudo.close()
    hojas = []
    for hoja in valores.worksheets:
        filas: dict[int, list[tuple[int, str]]] = {}
        n = 0
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.value is None or not hasattr(celda, "column"):
                    continue
                valor = texto_valor(celda.value).strip()
                if not valor:
                    continue
                n += 1
                if len(r.celdas) < MAX_CELDAS:
                    clave = (hoja.title, celda.row, celda.column)
                    r.celdas.append((hoja.title, celda.row, celda.column, valor,
                                     numero(celda.value) if not isinstance(celda.value, str) else None,
                                     formulas.get(clave)))
                    filas.setdefault(celda.row, []).append((celda.column, valor))
        filas_a_elementos(r, hoja.title, filas)
        hojas.append({"nombre": hoja.title, "dimension": hoja.calculate_dimension(), "celdas": n})
    props = valores.properties
    r.estructura = {
        "hojas": hojas,
        "formulas": len(formulas),
        "propiedades": {k: texto_valor(getattr(props, k)) for k in
                        ("title", "subject", "creator", "lastModifiedBy", "created", "modified")
                        if getattr(props, k, None)},
        "macros": zipfile.is_zipfile(path) and "xl/vbaProject.bin" in zipfile.ZipFile(path).namelist(),
    }
    valores.close()
    if len(r.celdas) >= MAX_CELDAS:
        r.mensaje = f"Recortado a {MAX_CELDAS} celdas."
    return r


def leer_ods(path: str) -> Resultado:
    import pandas as pd

    r = Resultado()
    libro = pd.read_excel(path, engine="odf", header=None, sheet_name=None)
    hojas = []
    for nombre, tabla in libro.items():
        filas: dict[int, list[tuple[int, str]]] = {}
        for (i, j), valor in tabla.stack().items():
            texto = texto_valor(valor).strip()
            if texto and len(r.celdas) < MAX_CELDAS:
                r.celdas.append((str(nombre), int(i) + 1, int(j) + 1, texto, numero(valor), None))
                filas.setdefault(int(i) + 1, []).append((int(j) + 1, texto))
        filas_a_elementos(r, str(nombre), filas)
        hojas.append({"nombre": str(nombre), "filas": int(tabla.shape[0]), "columnas": int(tabla.shape[1])})
    r.estructura = {"hojas": hojas}
    return r


# ---------------------------------------------------------------------------------------
# documentos: PPTX, DOCX, DOC, PDF, MSG, texto
# ---------------------------------------------------------------------------------------


def _formas(formas: Any) -> Any:
    for forma in formas:
        yield forma
        if getattr(forma, "shape_type", None) == 6:  # grupo
            yield from _formas(forma.shapes)


def leer_pptx(path: str) -> Resultado:
    from pptx import Presentation

    r = Resultado()
    presentacion = Presentation(path)
    total_imagenes = 0
    for numero_diapositiva, diapositiva in enumerate(presentacion.slides, 1):
        textos: list[str] = []
        imagenes = []
        for forma in _formas(diapositiva.shapes):
            if getattr(forma, "has_text_frame", False) and forma.has_text_frame:
                for parrafo in forma.text_frame.paragraphs:
                    t = "".join(run.text for run in parrafo.runs).strip() or parrafo.text.strip()
                    if t:
                        textos.append(t)
            if getattr(forma, "has_table", False) and forma.has_table:
                for fila in forma.table.rows:
                    celdas = [c.text.strip() for c in fila.cells]
                    if any(celdas):
                        textos.append(" | ".join(celdas))
            try:
                imagen = forma.image
            except (AttributeError, ValueError, KeyError):
                imagen = None
            if imagen is not None:
                imagenes.append({"fichero": imagen.filename, "bytes": len(imagen.blob),
                                 "sha1": imagen.sha1, "forma": forma.name})
        notas = ""
        if diapositiva.has_notes_slide:
            notas = diapositiva.notes_slide.notes_text_frame.text.strip()
        total_imagenes += len(imagenes)
        cuerpo = "\n".join(textos + ([f"[notas] {notas}"] if notas else []))
        r.elemento("diapositiva", f"diapositiva {numero_diapositiva}", cuerpo,
                   {"imagenes": imagenes} if imagenes else None)
    core = presentacion.core_properties
    r.estructura = {
        "diapositivas": len(presentacion.slides),
        "imagenes": total_imagenes,
        "propiedades": {k: texto_valor(getattr(core, k)) for k in
                        ("title", "author", "last_modified_by", "created", "modified")
                        if getattr(core, k, None)},
    }
    return r


def leer_docx(path: str) -> Resultado:
    import docx

    r = Resultado()
    documento = docx.Document(path)
    for i, parrafo in enumerate(documento.paragraphs, 1):
        if parrafo.text.strip():
            estilo = parrafo.style.name if parrafo.style is not None else ""
            r.elemento("parrafo", f"parrafo {i}", parrafo.text.strip(),
                       {"estilo": estilo} if estilo.lower().startswith(("heading", "título", "titulo")) else None)
    for t, tabla in enumerate(documento.tables, 1):
        for f, fila in enumerate(tabla.rows, 1):
            celdas = [c.text.strip() for c in fila.cells]
            if any(celdas):
                r.elemento("tabla", f"tabla {t} fila {f}", " | ".join(celdas))
    imagenes = [p for p in documento.part.package.parts if p.content_type.startswith("image/")]
    core = documento.core_properties
    r.estructura = {
        "parrafos": len(documento.paragraphs), "tablas": len(documento.tables), "imagenes": len(imagenes),
        "propiedades": {k: texto_valor(getattr(core, k)) for k in
                        ("title", "author", "last_modified_by", "created", "modified")
                        if getattr(core, k, None)},
    }
    if not r.elementos:
        r.estado = "sin_texto"
    return r


def leer_doc(path: str) -> Resultado:
    """Word 97-2003 por COM; si la directiva de Word bloquea el formato, texto legible del binario."""
    try:
        return _leer_doc_com(path)
    except Exception as error:  # pywintypes.com_error: formato bloqueado por directiva
        # Codigo -2146821984: formato bloqueado en el Centro de confianza ("se bloqueó", "blocked")
        if "-2146821984" not in str(error) and "bloque" not in str(error).lower() and "block" not in str(error).lower():
            raise
    r = leer_binario(path, "Word antiguo bloqueado por la directiva de Word: texto aproximado del binario.",
                     cabecera=16 * 1024 * 1024)
    r.estado = "extraido" if r.elementos else "sin_texto"
    r.estructura["propiedades"] = propiedades_ole(path)
    return r


def _leer_doc_com(path: str) -> Resultado:
    import pythoncom
    import win32com.client

    r = Resultado()
    carpeta, copia = copia_temporal(path)
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        documento = word.Documents.Open(copia, False, True, False)  # ConfirmConversions, ReadOnly, AddToRecent
        try:
            for i in range(1, documento.Paragraphs.Count + 1):
                t = documento.Paragraphs(i).Range.Text.strip()
                if t:
                    r.elemento("parrafo", f"parrafo {i}", t)
            r.estructura = {"parrafos": documento.Paragraphs.Count, "tablas": documento.Tables.Count,
                            "imagenes": documento.InlineShapes.Count + documento.Shapes.Count,
                            "propiedades": propiedades_ole(path)}
        finally:
            documento.Close(False)
    finally:
        word.Quit()
        pythoncom.CoUninitialize()
        shutil.rmtree(carpeta, ignore_errors=True)
    if not r.elementos:
        r.estado = "sin_texto"
    return r


def leer_pdf(path: str) -> Resultado:
    import fitz

    r = Resultado()
    documento = fitz.open(path)
    caracteres = 0
    imagenes = 0
    for i, pagina in enumerate(documento, 1):
        texto = pagina.get_text().strip()
        caracteres += len(texto)
        n_img = len(pagina.get_images(full=True))
        imagenes += n_img
        r.elemento("pagina", f"pagina {i}", texto,
                   {"imagenes": n_img, "ancho_pt": round(pagina.rect.width), "alto_pt": round(pagina.rect.height)})
    r.estructura = {"paginas": documento.page_count, "caracteres": caracteres, "imagenes": imagenes,
                    "metadatos": {k: v for k, v in (documento.metadata or {}).items() if v},
                    "fuentes": len({f[3] for p in documento for f in p.get_fonts()})}
    documento.close()
    if caracteres < 20:
        r.estado = "sin_texto"
        r.mensaje = "PDF sin capa de texto (escaneo o solo imagen)."
    return r


def leer_msg(path: str) -> Resultado:
    import extract_msg

    r = Resultado()
    mensaje = extract_msg.openMsg(path)
    try:
        adjuntos = [{"nombre": a.longFilename or a.shortFilename,
                     "bytes": len(a.data) if isinstance(a.data, bytes) else None}
                    for a in mensaje.attachments]
        cabecera = f"Asunto: {mensaje.subject}\nDe: {mensaje.sender}\nPara: {mensaje.to}\nCC: {mensaje.cc}\nFecha: {mensaje.date}"
        r.elemento("cabecera", "cabecera", cabecera)
        cuerpo = (mensaje.body or "").strip()
        for k, bloque in enumerate(re.split(r"\n{3,}", cuerpo), 1):
            if bloque.strip():
                r.elemento("cuerpo", f"cuerpo {k}", bloque.strip())
        r.estructura = {"asunto": mensaje.subject, "fecha": texto_valor(mensaje.date),
                        "remitente": mensaje.sender, "adjuntos": adjuntos}
    finally:
        mensaje.close()
    return r


def es_nube_de_puntos(lineas: list[str]) -> bool:
    muestra = [ln.split() for ln in lineas[:300] if ln.strip()]
    if len(muestra) < 20:
        return False
    validas = sum(1 for p in muestra if 2 <= len(p) <= 6 and all(numero(x) is not None for x in p))
    return validas / len(muestra) > 0.95


def leer_texto(path: str) -> Resultado:
    r = Resultado()
    texto, codificacion = decodificar(Path(path).read_bytes())
    lineas = texto.splitlines()
    if es_nube_de_puntos(lineas):
        import numpy as np

        puntos = [[numero(x) for x in ln.split()] for ln in lineas if ln.strip()]
        columnas = {len(p) for p in puntos}
        estructura: dict[str, Any] = {"tipo": "nube_de_puntos", "puntos": len(puntos),
                                      "columnas": sorted(columnas), "codificacion": codificacion}
        if len(columnas) == 1:
            matriz = np.array(puntos, dtype=float)
            estructura["minimos"] = [round(float(v), 4) for v in matriz.min(axis=0)]
            estructura["maximos"] = [round(float(v), 4) for v in matriz.max(axis=0)]
        estructura["primeras_lineas"] = lineas[:3]
        r.estructura = estructura
        r.elemento("resumen", "resumen", f"Nube de puntos: {len(puntos)} filas, columnas {sorted(columnas)}",
                   estructura)
        return r
    bloque: list[str] = []
    inicio = 1
    for i, linea in enumerate(lineas[:MAX_LINEAS_TEXTO], 1):
        bloque.append(linea)
        if len(bloque) == 40:
            if any(x.strip() for x in bloque):
                r.elemento("lineas", f"lineas {inicio}-{i}", "\n".join(bloque).strip())
            bloque, inicio = [], i + 1
    if any(x.strip() for x in bloque):
        r.elemento("lineas", f"lineas {inicio}-{inicio + len(bloque) - 1}", "\n".join(bloque).strip())
    r.estructura = {"tipo": "texto", "lineas": len(lineas), "codificacion": codificacion}
    if len(lineas) > MAX_LINEAS_TEXTO:
        r.mensaje = f"Recortado a {MAX_LINEAS_TEXTO} lineas."
    if not r.elementos:
        r.estado = "sin_texto"
    return r


# ---------------------------------------------------------------------------------------
# CAD y mallas
# ---------------------------------------------------------------------------------------

STEP_ENTIDAD = re.compile(rb"#\d+\s*=\s*([A-Z][A-Z0-9_]*)\s*\(")
STEP_PUNTO = re.compile(
    rb"#(\d+)\s*=\s*CARTESIAN_POINT\s*\(\s*'[^']*'\s*,\s*\(\s*([-+0-9.Ee]+)\s*,\s*([-+0-9.Ee]+)\s*,\s*([-+0-9.Ee]+)\s*\)"
)
STEP_VERTICE = re.compile(rb"VERTEX_POINT\s*\(\s*'[^']*'\s*,\s*#(\d+)\s*\)")
STEP_PRODUCTO = re.compile(rb"=\s*PRODUCT\s*\(\s*'((?:[^']|'')*)'\s*,\s*'((?:[^']|'')*)'")


def leer_step(path: str) -> Resultado:
    """Cabecera, productos, recuento de entidades y caja de los vertices de la geometria.

    La caja usa solo los CARTESIAN_POINT referenciados por VERTEX_POINT: los demas puntos
    (origenes de ejes, planos auxiliares) pueden estar muy lejos de la pieza.
    """
    from array import array

    import numpy as np

    r = Resultado()
    recuento: dict[str, int] = {}
    productos: dict[str, None] = {}
    ids_punto = array("q")
    coords = array("d")
    ids_vertice = array("q")
    cabecera = b""
    unidades = set()
    resto = b""
    with open(path, "rb") as f:
        primera = True
        while True:
            trozo = f.read(32 * 1024 * 1024)
            if not trozo:
                break
            datos = resto + trozo
            corte = datos.rfind(b";")
            if corte < 0:
                resto = datos
                continue
            datos, resto = datos[: corte + 1], datos[corte + 1:]
            if primera:
                fin = datos.find(b"ENDSEC;")
                cabecera = datos[: fin + 7] if fin >= 0 else datos[:4000]
                primera = False
            for m in STEP_ENTIDAD.finditer(datos):
                nombre = m.group(1).decode()
                recuento[nombre] = recuento.get(nombre, 0) + 1
            for m in STEP_PUNTO.finditer(datos):
                xyz = [numero(m.group(k).decode()) for k in (2, 3, 4)]
                if None not in xyz:
                    ids_punto.append(int(m.group(1)))
                    coords.extend(xyz)  # type: ignore[arg-type]
            ids_vertice.extend(int(m.group(1)) for m in STEP_VERTICE.finditer(datos))
            for m in STEP_PRODUCTO.finditer(datos):
                if len(productos) < 2000:
                    productos.setdefault(m.group(1).decode("latin-1"), None)
            for u in re.findall(rb"SI_UNIT\s*\(\s*\.?([A-Z]*)\.?\s*,\s*\.([A-Z]+)\.", datos):
                unidades.add(b" ".join(u).decode())
    caja = None
    if len(ids_punto):
        todos = np.frombuffer(ids_punto, dtype=np.int64)
        puntos = np.frombuffer(coords, dtype=np.float64).reshape(-1, 3)
        usados = np.isin(todos, np.frombuffer(ids_vertice, dtype=np.int64)) if len(ids_vertice) else None
        seleccion = puntos[usados] if usados is not None and usados.any() else puntos
        minimo, maximo = seleccion.min(axis=0), seleccion.max(axis=0)
        caja = {"min": [round(float(v), 3) for v in minimo], "max": [round(float(v), 3) for v in maximo],
                "dimensiones": [round(float(b - a), 3) for a, b in zip(minimo, maximo)],
                "base": "vertices" if usados is not None and usados.any() else "todos los puntos"}
    cab = cabecera.decode("latin-1", "ignore")
    r.elemento("cabecera", "cabecera", cab.strip())
    if productos:
        r.elemento("productos", "productos", "\n".join(productos))
    claves = ["PRODUCT", "MANIFOLD_SOLID_BREP", "BREP_WITH_VOIDS", "CLOSED_SHELL", "OPEN_SHELL",
              "ADVANCED_FACE", "CYLINDRICAL_SURFACE", "PLANE", "B_SPLINE_SURFACE_WITH_KNOTS",
              "CONICAL_SURFACE", "TOROIDAL_SURFACE", "SPHERICAL_SURFACE", "CARTESIAN_POINT",
              "NEXT_ASSEMBLY_USAGE_OCCURRENCE", "SHELL_BASED_SURFACE_MODEL", "STYLED_ITEM"]
    r.estructura = {
        "sistema_origen": (re.search(r"FILE_NAME\s*\((.*?)\);", cab, re.S) or [None, None])[1],
        "esquema": (re.search(r"FILE_SCHEMA\s*\(\s*\(\s*'([^']*)'", cab) or [None, None])[1],
        "productos": len(productos),
        "entidades": {k: recuento.get(k, 0) for k in claves if recuento.get(k)},
        "entidades_total": sum(recuento.values()),
        "unidades": sorted(unidades),
        "caja": caja,
    }
    return r


def leer_stl(path: str) -> Resultado:
    import numpy as np

    r = Resultado()
    tamano = os.path.getsize(path)
    with open(path, "rb") as f:
        cabecera = f.read(84)
    n = int.from_bytes(cabecera[80:84], "little") if len(cabecera) == 84 else 0
    if tamano == 84 + 50 * n and n:
        tipo = np.dtype([("n", "<f4", 3), ("v", "<f4", (3, 3)), ("a", "<u2")])
        mallas = np.memmap(path, dtype=tipo, mode="r", offset=84, shape=(n,))
        v = mallas["v"].reshape(-1, 3)
        minimo, maximo = v.min(axis=0), v.max(axis=0)
        formato = "binario"
    else:
        vertices = []
        with open(path, "rb") as f:
            for linea in f:
                linea = linea.strip()
                if linea.startswith(b"vertex"):
                    vertices.append([float(x) for x in linea.split()[1:4]])
        v = np.array(vertices) if vertices else np.zeros((0, 3))
        n = len(vertices) // 3
        minimo, maximo = (v.min(axis=0), v.max(axis=0)) if len(v) else (np.zeros(3), np.zeros(3))
        formato = "ascii"
    r.estructura = {
        "formato": formato, "triangulos": n,
        "cabecera": cabecera[:80].decode("latin-1", "ignore").strip("\x00 "),
        "caja": {"min": [round(float(x), 3) for x in minimo], "max": [round(float(x), 3) for x in maximo]},
        "dimensiones": [round(float(b - a), 3) for a, b in zip(minimo, maximo)],
    }
    r.elemento("resumen", "resumen", json.dumps(r.estructura, ensure_ascii=False))
    return r


def leer_iges(path: str) -> Resultado:
    r = Resultado()
    secciones: dict[str, list[str]] = {"S": [], "G": [], "D": []}
    with open(path, "rb") as f:
        for linea in f:
            texto = linea.decode("latin-1").rstrip("\r\n")
            if len(texto) >= 73 and texto[72] in secciones:
                secciones[texto[72]].append(texto[:72])
    tipos: dict[str, int] = {}
    for i in range(0, len(secciones["D"]), 2):
        tipo = secciones["D"][i][:8].strip()
        tipos[tipo] = tipos.get(tipo, 0) + 1
    inicio = "\n".join(s.rstrip() for s in secciones["S"]).strip()
    global_ = "".join(secciones["G"]).strip()
    r.elemento("cabecera", "seccion S", inicio)
    r.elemento("cabecera", "seccion G", global_)
    r.estructura = {"entidades_por_tipo": tipos, "entidades": sum(tipos.values())}
    return r


def leer_dxf(path: str) -> Resultado:
    import ezdxf
    from ezdxf import recover

    r = Resultado()
    documento, _ = recover.readfile(path)
    espacio = documento.modelspace()
    tipos: dict[str, int] = {}
    textos = []
    vertices = 0
    caja = [math.inf, math.inf, math.inf, -math.inf, -math.inf, -math.inf]
    for entidad in espacio:
        tipos[entidad.dxftype()] = tipos.get(entidad.dxftype(), 0) + 1
        if entidad.dxftype() in ("TEXT", "MTEXT") and len(textos) < 2000:
            textos.append(entidad.plain_text() if entidad.dxftype() == "MTEXT" else entidad.dxf.text)
        # Los perfiles de la CMM son una sola POLYLINE 3D sin texto: lo util es cuantos puntos y donde.
        if entidad.dxftype() in ("POLYLINE", "LWPOLYLINE"):
            puntos = ([tuple(v.dxf.location) for v in entidad.vertices] if entidad.dxftype() == "POLYLINE"
                      else [(x, y, 0.0) for x, y, *_ in entidad.get_points()])
            vertices += len(puntos)
            for punto in puntos:
                for i, valor in enumerate(punto[:3]):
                    caja[i], caja[i + 3] = min(caja[i], valor), max(caja[i + 3], valor)
    if textos:
        r.elemento("textos", "textos del modelo", "\n".join(textos))
    r.estructura = {"version": documento.dxfversion, "unidades": documento.header.get("$INSUNITS"),
                    "entidades": tipos, "capas": [c.dxf.name for c in documento.layers][:200],
                    "ezdxf": ezdxf.__version__}
    if vertices:
        r.estructura["vertices"] = vertices
        r.estructura["caja"] = [round(v, 4) for v in caja]
    return r


def leer_binario(path: str, nota: str, cabecera: int = 4 * 1024 * 1024) -> Resultado:
    """Formatos propietarios (CATIA, Moldflow, DWG...): firma y cadenas legibles."""
    r = Resultado(estado="solo_metadatos", mensaje=nota)
    tamano = os.path.getsize(path)
    with open(path, "rb") as f:
        inicio = f.read(cabecera)
        f.seek(max(0, tamano - cabecera))
        final = f.read(cabecera) if tamano > cabecera else b""
    cadenas = cadenas_imprimibles(inicio + b"\n" + final)
    r.estructura = {"firma_hex": inicio[:16].hex(), "firma_texto": inicio[:16].decode("latin-1", "ignore"),
                    "cadenas": len(cadenas)}
    if cadenas:
        r.elemento("cadenas", "cadenas legibles (inicio y final)", "\n".join(cadenas))
    return r


def leer_catia(path: str) -> Resultado:
    r = leer_binario(path, "CATIA V5 nativo: sin lector; solo firma y cadenas legibles.")
    texto = "\n".join(e[2] for e in r.elementos)
    version = re.findall(r"V5R\d+|R\d{2} SP\d+|CATIA V\d[^\n]{0,30}", texto)
    r.estructura["version_catia"] = sorted(set(version))[:10]
    return r


# ---------------------------------------------------------------------------------------
# multimedia y archivos comprimidos
# ---------------------------------------------------------------------------------------


def leer_imagen(path: str) -> Resultado:
    from PIL import ExifTags, Image

    r = Resultado(estado="solo_metadatos", mensaje="Imagen: revisar visualmente si es relevante.")
    with Image.open(path) as imagen:
        exif = {}
        try:
            for k, v in (imagen.getexif() or {}).items():
                nombre = ExifTags.TAGS.get(k, str(k))
                if nombre in ("DateTime", "DateTimeOriginal", "Model", "Make", "Software"):
                    exif[nombre] = texto_valor(v)
        except Exception:
            pass
        r.estructura = {"formato": imagen.format, "ancho": imagen.width, "alto": imagen.height,
                        "modo": imagen.mode, "exif": exif}
    return r


def leer_video(path: str) -> Resultado:
    import cv2

    r = Resultado(estado="solo_metadatos", mensaje="Video: revisar fotogramas si es relevante.")
    captura = cv2.VideoCapture(path)
    try:
        fps = captura.get(cv2.CAP_PROP_FPS) or 0
        fotogramas = captura.get(cv2.CAP_PROP_FRAME_COUNT) or 0
        r.estructura = {"ancho": int(captura.get(cv2.CAP_PROP_FRAME_WIDTH)),
                        "alto": int(captura.get(cv2.CAP_PROP_FRAME_HEIGHT)),
                        "fps": round(fps, 2), "fotogramas": int(fotogramas),
                        "duracion_s": round(fotogramas / fps, 1) if fps else None}
    finally:
        captura.release()
    return r


def leer_comprimido(path: str) -> Resultado:
    """Lista el contenido. Su extraccion e inventario se hace aparte (ver AVANCE.md)."""
    r = Resultado()
    entradas = []
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            for info in z.infolist():
                if not info.is_dir():
                    entradas.append({"ruta": info.filename, "bytes": info.file_size,
                                     "fecha": "%04d-%02d-%02d %02d:%02d" % info.date_time[:5]})
    else:
        salida = subprocess.run([SIETE_ZIP, "l", "-slt", "-ba", path], capture_output=True,
                                text=True, encoding="utf-8", errors="replace", timeout=300)
        actual: dict[str, str] = {}
        for linea in salida.stdout.splitlines() + [""]:
            if not linea.strip():
                if actual.get("Path") and actual.get("Folder") != "+":
                    entradas.append({"ruta": actual["Path"], "bytes": int(actual.get("Size") or 0),
                                     "fecha": actual.get("Modified", "")})
                actual = {}
            elif " = " in linea:
                k, v = linea.split(" = ", 1)
                actual[k.strip()] = v.strip()
        if salida.returncode and not entradas:
            raise ValueError(salida.stderr.strip()[:300] or "7-Zip no puede listar el archivo")
    r.estructura = {"entradas": len(entradas), "bytes": sum(e["bytes"] for e in entradas)}
    r.elemento("entradas", "contenido", "\n".join(f"{e['ruta']} ({e['bytes']} B, {e['fecha']})"
                                                     for e in entradas), {"entradas": entradas[:2000]})
    return r


def leer_acceso_directo(path: str) -> Resultado:
    import pythoncom
    import win32com.client

    r = Resultado()
    pythoncom.CoInitialize()
    try:
        acceso = win32com.client.Dispatch("WScript.Shell").CreateShortcut(path.removeprefix("\\\\?\\"))
        r.estructura = {"destino": acceso.TargetPath, "argumentos": acceso.Arguments}
        r.elemento("destino", "destino", acceso.TargetPath)
    finally:
        pythoncom.CoUninitialize()
    return r


def leer_generico(path: str) -> Resultado:
    """Texto si lo parece; si no, firma y cadenas legibles."""
    with open(path, "rb") as f:
        inicio = f.read(65536)
    if inicio and sum(32 <= b < 127 or b in (9, 10, 13) for b in inicio) / len(inicio) > 0.9:
        r = leer_texto(path)
        r.estructura["deteccion"] = "texto"
        return r
    return leer_binario(path, "Formato binario sin lector especifico.")


def ignorar(nota: str) -> Callable[[str], Resultado]:
    def lector(path: str) -> Resultado:
        return Resultado(estado="solo_metadatos", mensaje=nota)
    return lector


# extension -> (nombre del lector, version, funcion). Subir la version obliga a re-extraer.
LECTORES: dict[str, tuple[str, int, Callable[[str], Resultado]]] = {
    "csv": ("csv", 1, leer_csv),
    "xls": ("xls", 1, leer_xls),
    "xlsx": ("xlsx", 1, leer_xlsx),
    "xlsm": ("xlsx", 1, leer_xlsx),
    "ods": ("ods", 1, leer_ods),
    "pptx": ("pptx", 1, leer_pptx),
    "docx": ("docx", 1, leer_docx),
    "doc": ("doc-com", 1, leer_doc),
    "pdf": ("pdf", 1, leer_pdf),
    "msg": ("msg", 1, leer_msg),
    "txt": ("texto", 1, leer_texto),
    "log": ("texto", 1, leer_texto),
    "stp": ("step", 2, leer_step),
    "step": ("step", 2, leer_step),
    "stl": ("stl", 1, leer_stl),
    "igs": ("iges", 1, leer_iges),
    "iges": ("iges", 1, leer_iges),
    "dxf": ("dxf", 2, leer_dxf),
    "catpart": ("catia", 1, leer_catia),
    "catproduct": ("catia", 1, leer_catia),
    "catdrawing": ("catia", 1, leer_catia),
    "mfr": ("binario", 1, lambda p: leer_binario(p, "Moldflow (.mfr): formato propietario sin lector.")),
    "dwg": ("binario", 1, lambda p: leer_binario(p, "DWG: formato propietario sin lector.")),
    "jpg": ("imagen", 1, leer_imagen),
    "jpeg": ("imagen", 1, leer_imagen),
    "jfif": ("imagen", 1, leer_imagen),
    "png": ("imagen", 1, leer_imagen),
    "tif": ("imagen", 1, leer_imagen),
    "mov": ("video", 1, leer_video),
    "mp4": ("video", 1, leer_video),
    "avi": ("video", 1, leer_video),
    "zip": ("comprimido", 1, leer_comprimido),
    "rar": ("comprimido", 1, leer_comprimido),
    "cab": ("comprimido", 1, leer_comprimido),
    "lnk": ("acceso-directo", 1, leer_acceso_directo),
}


def sin_celdas_es_sin_texto(leer: Callable[[str], Resultado]) -> Callable[[str], Resultado]:
    """Un libro sin ninguna celda no es 'extraido': su contenido esta en imagenes o graficos."""

    def envoltura(path: str) -> Resultado:
        r = leer(path)
        if r.estado == "extraido" and not r.celdas and not r.elementos:
            imagenes = None
            # Solo OOXML/ODF: un .xls OLE con objetos incrustados tambien "parece" un zip.
            if Path(path).suffix.lower() in (".xlsx", ".xlsm", ".ods") and zipfile.is_zipfile(path):
                with zipfile.ZipFile(path) as z:
                    imagenes = sum(1 for n in z.namelist() if re.search(r"(^|/)media/", n))
            r.estado = "sin_texto"
            r.mensaje = ("Libro sin celdas: el contenido esta en imagenes o graficos"
                         + (f" ({imagenes} imagenes)." if imagenes is not None else "."))
        return r

    return envoltura


def lector_para(nombre: str, ext: str, ruta: str) -> tuple[str, int, Callable[[str], Resultado]]:
    if nombre.startswith("~$"):
        return ("ignorado", 1, ignorar("Fichero de bloqueo de Office (~$): sin contenido del proyecto."))
    if nombre.lower() == "thumbs.db":
        return ("ignorado", 1, ignorar("Miniaturas de Windows (Thumbs.db): sin contenido del proyecto."))
    if "/gpucache/" in ruta.lower():
        return ("ignorado", 1, ignorar("Cache de GPU de una aplicacion: sin contenido del proyecto."))
    if ext in ("xls", "xlsx", "xlsm", "ods"):
        nombre_lector, version, leer = LECTORES[ext]
        return (nombre_lector, version, sin_celdas_es_sin_texto(leer))
    return LECTORES.get(ext, ("generico", 1, leer_generico))
